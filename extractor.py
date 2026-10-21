"""Estrae il testo dai documenti caricati: PDF, Word ed Excel."""

import docx
import openpyxl
from pypdf import PdfReader


def _pdf(file):
    reader = PdfReader(file)
    pagine = []
    for i, pagina in enumerate(reader.pages, 1):
        pagine.append(f"[Pagina {i}]\n{pagina.extract_text() or ''}")
    return "\n\n".join(pagine)


def _docx(file):
    documento = docx.Document(file)
    righe = [p.text for p in documento.paragraphs if p.text.strip()]
    for tabella in documento.tables:
        for riga in tabella.rows:
            righe.append(" | ".join(c.text.strip() for c in riga.cells))
    return "\n".join(righe)


def _xlsx(file):
    libro = openpyxl.load_workbook(file, data_only=True, read_only=True)
    parti = []
    for foglio in libro.worksheets:
        righe = []
        for riga in foglio.iter_rows(values_only=True):
            if any(v is not None for v in riga):
                righe.append(" | ".join("" if v is None else str(v) for v in riga))
        parti.append(f"[Foglio: {foglio.title}]\n" + "\n".join(righe))
    return "\n\n".join(parti)


def extract_text(file):
    """Riceve un file caricato da Streamlit (ha l'attributo .name) e restituisce il testo."""
    nome = file.name.lower()
    if nome.endswith(".pdf"):
        return _pdf(file)
    if nome.endswith(".docx"):
        return _docx(file)
    if nome.endswith(".xlsx"):
        return _xlsx(file)
    return ""


def build_knowledge_base(files):
    """Unisce il testo di tutti i file, ciascuno con il suo nome come intestazione."""
    return "\n\n".join(f"===== DOCUMENTO: {f.name} =====\n{extract_text(f)}" for f in files)
