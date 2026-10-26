"""Generazione di file (PDF, Word, Excel) e grafici, sempre in memoria con BytesIO."""

import json
import re
import textwrap
from datetime import date
from io import BytesIO

import docx
import matplotlib
import openpyxl
from docx.shared import Inches
from fpdf import FPDF, FontFace
from openpyxl.styles import Font, PatternFill

matplotlib.use("Agg")  # nessuna finestra: il grafico viene solo salvato in memoria
import matplotlib.pyplot as plt  # noqa: E402

VERDE = "1F4E5A"
MIME = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "json": "application/json",
}


# ---------- lettura del contenuto in markdown semplice ----------

def _blocchi(testo):
    """Divide il testo in blocchi: titoli, elenchi, tabelle, paragrafi."""
    blocchi, tabella = [], []
    for riga in testo.splitlines() + [""]:
        r = riga.strip()
        if "|" in r:
            celle = [c.strip() for c in r.strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in celle if c):  # salta |---|---|
                tabella.append(celle)
            continue
        if tabella:
            blocchi.append(("tabella", tabella))
            tabella = []
        if not r or re.fullmatch(r"-{3,}|\*{3,}|_{3,}", r):
            continue
        if r.startswith("> "):
            r = r[2:].strip()
        if r.startswith("### "):
            blocchi.append(("h3", r[4:]))
        elif r.startswith("## "):
            blocchi.append(("h2", r[3:]))
        elif r.startswith("# "):
            blocchi.append(("h1", r[2:]))
        elif r.startswith(("- ", "* ")):
            blocchi.append(("punto", r[2:]))
        else:
            blocchi.append(("testo", r))
    return blocchi


def _pulisci(testo):
    """Toglie la formattazione markdown (grassetto, corsivo, codice)."""
    testo = testo.replace("**", "").replace("`", "")
    return re.sub(r"(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*])", r"\1", testo)


def _numero(valore):
    """Converte '1.234,5' o '1234.5' o '12%' in numero, se possibile."""
    s = valore.replace(" ", "").replace(" ", "")
    perc = s.endswith("%")
    s = s.rstrip("%")
    if re.fullmatch(r"-?\d{1,3}(\.\d{3})+(,\d+)?", s) or re.fullmatch(r"-?\d+,\d+", s):
        s = s.replace(".", "").replace(",", ".")
    try:
        n = float(s)
    except ValueError:
        return valore
    return n / 100 if perc else n


# ---------- grafico ----------

def make_chart(titolo, etichette, valori, unita, serie_confronto=None, nomi_serie=None):
    fig, ax = plt.subplots(figsize=(8, 4.5), dpi=150)
    x = range(len(etichette))
    if serie_confronto and len(serie_confronto) == len(valori):
        nomi = nomi_serie if nomi_serie and len(nomi_serie) == 2 else ["Serie 1", "Serie 2"]
        w = 0.38
        b1 = ax.bar([i - w / 2 for i in x], valori, w, label=nomi[0], color="#1F4E5A")
        b2 = ax.bar([i + w / 2 for i in x], serie_confronto, w, label=nomi[1], color="#9BC4BC")
        ax.legend(frameon=False)
        barre = list(b1) + list(b2)
    else:
        barre = ax.bar(list(x), valori, color="#1F4E5A")
    for b in barre:
        ax.annotate(f"{b.get_height():,.1f}".replace(",", "X").replace(".", ",").replace("X", "."),
                    (b.get_x() + b.get_width() / 2, b.get_height()), ha="center", va="bottom", fontsize=8)
    ax.set_xticks(list(x), [textwrap.fill(str(e), 14) for e in etichette], fontsize=9)
    ax.set_ylabel(unita)
    ax.set_title(titolo, fontsize=12, fontweight="bold", color="#1F4E5A")
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    buf = BytesIO()
    fig.savefig(buf, format="png")
    plt.close(fig)
    return buf.getvalue()


# ---------- Excel ----------

def _xlsx(titolo, contenuto):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.append([titolo])
    ws["A1"].font = Font(bold=True, size=14, color=VERDE)
    ws.append([])
    for tipo, valore in _blocchi(contenuto):
        if tipo == "tabella":
            for i, riga in enumerate(valore):
                if i == 0:
                    ws.append([_pulisci(c) for c in riga])
                    for cella in ws[ws.max_row]:
                        cella.font = Font(bold=True, color="FFFFFF")
                        cella.fill = PatternFill("solid", fgColor=VERDE)
                else:
                    ws.append([_numero(_pulisci(c)) for c in riga])
                    for cella, testo in zip(ws[ws.max_row], riga):
                        if isinstance(cella.value, float) and testo.strip().endswith("%"):
                            cella.number_format = "0.0%"
                        elif isinstance(cella.value, (int, float)):
                            cella.number_format = "#,##0.0"
            ws.append([])
        elif tipo in ("h1", "h2", "h3"):
            ws.append([_pulisci(valore)])
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12 if tipo == "h1" else 11, color=VERDE)
        else:
            ws.append([("• " if tipo == "punto" else "") + _pulisci(valore)])
    for colonna in ws.columns:
        lunghezza = max(len(str(c.value)) if c.value is not None and c.row > 2 else 0 for c in colonna)
        ws.column_dimensions[colonna[0].column_letter].width = min(max(12, lunghezza + 2), 60)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- Word ----------

def _docx(titolo, contenuto, grafico, foto=None):
    d = docx.Document()
    d.add_heading(titolo, 0)
    for tipo, valore in _blocchi(contenuto):
        if tipo == "h1":
            d.add_heading(_pulisci(valore), 1)
        elif tipo == "h2":
            d.add_heading(_pulisci(valore), 2)
        elif tipo == "h3":
            d.add_heading(_pulisci(valore), 3)
        elif tipo == "punto":
            d.add_paragraph(_pulisci(valore), style="List Bullet")
        elif tipo == "tabella":
            colonne = max(len(r) for r in valore)
            t = d.add_table(rows=0, cols=colonne)
            t.style = "Light Grid Accent 1"
            for riga in valore:
                celle = t.add_row().cells
                for i, c in enumerate(riga):
                    celle[i].text = _pulisci(c)
        else:
            d.add_paragraph(_pulisci(valore))
    if grafico:
        d.add_picture(BytesIO(grafico), width=Inches(6))
    if foto:
        d.add_heading("Documentazione fotografica", 2)
        for i, f in enumerate(foto, 1):
            d.add_picture(BytesIO(foto_pronta(f)[0]), width=Inches(4.5))
            d.add_paragraph(f"Foto {i}: fornita dall'utente, non verificata sul posto")
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


# ---------- PDF ----------

_SOSTITUZIONI = {"€": "EUR", "’": "'", "‘": "'", "“": '"', "”": '"', "–": "-", "—": "-", "•": "-",
                 "₂": "2", "³": "3", "…": "...", "≈": "~", "→": "->", "≥": ">=", "≤": "<=",
                 "⚠️": "(!)", "⚠": "(!)", "✅": "[OK]", "✔": "[OK]", "❌": "[X]", "✗": "[X]", "−": "-"}


def _latin1(testo):
    """I font standard del PDF usano latin-1: sostituiamo i caratteri che non ci stanno."""
    for a, b in _SOSTITUZIONI.items():
        testo = testo.replace(a, b)
    return testo.encode("latin-1", "ignore").decode("latin-1")  # le emoji rimaste spariscono


def _larghezze_colonne(righe):
    """Larghezze relative delle colonne in base al testo più lungo (con minimo e massimo)."""
    if righe and [c.strip() for c in righe[0]] == ["Campo", "Valore", "Fonte", "Stato"]:
        return (22, 30, 30, 22)  # tabelle del passaporto: stesse colonne in tutto il documento
    colonne = max(len(r) for r in righe)
    lung = [max(len(r[i]) if i < len(r) else 0 for r in righe) for i in range(colonne)]
    # la parola più lunga di ogni colonna non deve andare a capo (es. "documentato", nomi di file)
    parola = [max((len(w) for r in righe if i < len(r) for w in r[i].split()), default=0) for i in range(colonne)]
    return tuple(max(min(max(n, 6), 45), parola[i] + 2) for i, n in enumerate(lung))


# Palette "Ardesia", la stessa dell'app
ARDESIA, LICHENE, AMBRA = (22, 35, 42), (79, 140, 90), (210, 140, 40)
TESTO, GRIGIO, FILO, RIGA_ALTERNA = (38, 48, 52), (120, 134, 128), (208, 216, 212), (243, 246, 244)


def foto_pronta(dati, lato_max=1600):
    """Raddrizza la foto del telefono (orientamento EXIF), la rimpicciolisce e la restituisce in JPEG."""
    from PIL import Image, ImageOps
    img = ImageOps.exif_transpose(Image.open(BytesIO(dati))).convert("RGB")
    img.thumbnail((lato_max, lato_max))
    buf = BytesIO()
    img.save(buf, "JPEG", quality=85)
    return buf.getvalue(), img.width / img.height


class _Report(FPDF):
    """PDF con intestazione a fascia sulla prima pagina, testatina sulle altre e piè di pagina."""

    def __init__(self, titolo):
        super().__init__()
        self.titolo = _latin1(titolo)
        self.set_margins(18, 16, 18)
        self.set_auto_page_break(auto=True, margin=18)
        self.alias_nb_pages()

    def header(self):
        if self.page_no() == 1:
            self.set_fill_color(*ARDESIA)
            self.rect(0, 0, self.w, 38, "F")
            self.set_fill_color(*LICHENE)
            self.rect(0, 38, self.w, 1.6, "F")
            self.set_fill_color(*AMBRA)
            self.rect(0, 38, 34, 1.6, "F")
            self.set_xy(18, 10)
            self.set_text_color(255, 255, 255)
            self.set_font("Helvetica", "B", 18)
            self.multi_cell(self.w - 36, 8, self.titolo, new_x="LMARGIN", new_y="NEXT", align="L")
            self.set_x(18)
            self.set_text_color(170, 200, 180)
            self.set_font("Helvetica", size=8.5)
            self.cell(0, 6, f"Sustainability Copilot - generato il {date.today():%d/%m/%Y}")
            self.set_y(48)
        else:
            self.set_y(8)
            self.set_text_color(*GRIGIO)
            self.set_font("Helvetica", size=7.5)
            self.cell(0, 5, self.titolo)
            self.set_draw_color(*LICHENE)
            self.set_line_width(0.4)
            self.line(18, 14, self.w - 18, 14)
            self.set_y(20)

    def footer(self):
        self.set_y(-12)
        self.set_draw_color(*FILO)
        self.set_line_width(0.2)
        self.line(18, self.h - 14, self.w - 18, self.h - 14)
        self.set_text_color(*GRIGIO)
        self.set_font("Helvetica", size=7.5)
        self.cell(0, 6, "Documento generato con l'AI: da verificare prima dell'uso")
        self.set_x(18)
        self.cell(0, 6, f"Pagina {self.page_no()} di {{nb}}", align="R")


_ETICHETTA = re.compile(r"^([A-Za-zÀ-ÿ'][A-Za-zÀ-ÿ' ()/.]{1,38}):\s*(.*)$")


def _riga_testo(pdf, testo):
    """Paragrafo; le righe 'Etichetta: valore' hanno l'etichetta in grassetto, quelle con ____ diventano firme."""
    pdf.set_text_color(*TESTO)
    m = _ETICHETTA.match(testo)
    if m and len(m.group(1).split()) <= 5:
        etichetta, valore = m.group(1), m.group(2)
        pdf.set_font("Helvetica", "B", 10)
        pdf.cell(pdf.get_string_width(etichetta + ": ") + 1, 5.6, etichetta + ":")
        if "___" in valore or not valore.strip():  # riga per la firma
            x, y = pdf.get_x() + 1, pdf.get_y() + 4.6
            pdf.set_draw_color(*GRIGIO)
            pdf.set_line_width(0.3)
            pdf.line(x, y, min(x + 80, pdf.w - 18), y)
            pdf.ln(9)
            return
        pdf.set_font("Helvetica", size=10)
        pdf.multi_cell(0, 5.6, valore, new_x="LMARGIN", new_y="NEXT", align="L")
    else:
        pdf.set_font("Helvetica", size=10)
        pdf.multi_cell(0, 5.6, testo, new_x="LMARGIN", new_y="NEXT", align="L")
    pdf.ln(1.2)


def _pdf(titolo, contenuto, grafico, foto=None):
    pdf = _Report(titolo)
    pdf.add_page()
    for tipo, valore in _blocchi(contenuto):
        testo = _latin1(_pulisci(valore)) if tipo != "tabella" else None
        if tipo in ("h1", "h2", "h3") and "_" in testo and " " not in testo:
            testo = testo.replace("_", " ").capitalize()  # es. dati_mancanti -> Dati mancanti
        if tipo == "h1" and testo.strip().lower() == pdf.titolo.strip().lower():
            continue  # il titolo è già nella fascia
        if tipo == "tabella":
            colonne = max(len(r) for r in valore)
            pdf.set_font("Helvetica", size=8.5 if colonne <= 4 else 7.5)  # tabelle larghe: font più piccolo
            pdf.set_text_color(*TESTO)
            pdf.set_draw_color(*FILO)
            pdf.set_line_width(0.2)
            righe = [[_latin1(_pulisci(c)) for c in r] + [""] * (colonne - len(r)) for r in valore]
            stile = FontFace(emphasis="BOLD", color=255, fill_color=ARDESIA)
            pdf.set_fill_color(*RIGA_ALTERNA)  # righe alterne grigio chiaro (la tabella usa il colore corrente)
            with pdf.table(first_row_as_headings=True, headings_style=stile, line_height=5.6, text_align="LEFT",
                           col_widths=_larghezze_colonne(righe), cell_fill_color=RIGA_ALTERNA,
                           cell_fill_mode="ROWS", borders_layout="HORIZONTAL_LINES") as tabella:
                for r in righe:
                    riga = tabella.row()
                    for c in r:
                        riga.cell(c)
            pdf.ln(4)
        elif tipo in ("h1", "h2"):
            pdf.ln(3)
            if pdf.get_y() > pdf.h - 40:  # un titolo non resta da solo in fondo alla pagina
                pdf.add_page()
            y = pdf.get_y()
            pdf.set_fill_color(*LICHENE)
            pdf.rect(18, y + 1, 2.2, 5.6, "F")
            pdf.set_x(23)
            pdf.set_text_color(*ARDESIA)
            pdf.set_font("Helvetica", "B", 13 if tipo == "h1" else 12)
            pdf.multi_cell(0, 7.5, testo, new_x="LMARGIN", new_y="NEXT", align="L")
            pdf.ln(1.5)
        elif tipo == "h3":
            pdf.ln(1.5)
            if pdf.get_y() > pdf.h - 40:
                pdf.add_page()
            pdf.set_text_color(*LICHENE)
            pdf.set_font("Helvetica", "B", 10.5)
            pdf.multi_cell(0, 6.5, testo, new_x="LMARGIN", new_y="NEXT", align="L")
        elif tipo == "punto":
            y = pdf.get_y()
            pdf.set_fill_color(*LICHENE)
            pdf.ellipse(19.2, y + 2.1, 1.7, 1.7, "F")
            pdf.set_x(24)
            pdf.set_text_color(*TESTO)
            pdf.set_font("Helvetica", size=10)
            pdf.multi_cell(0, 5.6, testo, new_x="LMARGIN", new_y="NEXT", align="L")
            pdf.ln(1)
        else:
            _riga_testo(pdf, testo)
    if grafico:
        pdf.ln(3)
        pdf.image(BytesIO(grafico), w=pdf.epw)
    for i, f in enumerate(foto or [], 1):
        dati, proporzione = foto_pronta(f)
        larghezza = min(pdf.epw, 120 * proporzione) if proporzione < 1 else min(pdf.epw, 150)
        altezza = larghezza / proporzione
        if pdf.get_y() + altezza + 16 > pdf.h - 18:
            pdf.add_page()
        if i == 1:
            pdf.ln(2)
            pdf.set_text_color(*ARDESIA)
            pdf.set_font("Helvetica", "B", 12)
            pdf.cell(0, 8, "Documentazione fotografica", new_x="LMARGIN", new_y="NEXT", align="L")
        x = 18 + (pdf.epw - larghezza) / 2
        pdf.image(BytesIO(dati), x=x, w=larghezza)
        pdf.set_text_color(*GRIGIO)
        pdf.set_font("Helvetica", "I", 8)
        pdf.cell(0, 6, f"Foto {i}: fornita dall'utente, non verificata sul posto", align="C",
                 new_x="LMARGIN", new_y="NEXT")
        pdf.ln(2)
    return bytes(pdf.output())


PRIORITA = {"alta": ((185, 50, 40), (251, 235, 235)), "media": ((176, 112, 10), (252, 243, 225)),
            "bassa": ((60, 120, 70), (229, 244, 228))}


def _box(pdf, x, y, w, h, colore):
    pdf.set_fill_color(*colore)
    pdf.rect(x, y, w, h, "F", round_corners=True, corner_radius=2)


def _titolo_sezione(pdf, testo):
    pdf.ln(3)
    if pdf.get_y() > pdf.h - 50:
        pdf.add_page()
    y = pdf.get_y()
    pdf.set_fill_color(*LICHENE)
    pdf.rect(18, y + 1, 2.2, 5.6, "F")
    pdf.set_xy(23, y)
    pdf.set_text_color(*ARDESIA)
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 7.5, _latin1(testo), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(1.5)


def _tabella(pdf, intestazione, righe, larghezze, colonna_priorita=None):
    pdf.set_font("Helvetica", size=8.8)
    pdf.set_text_color(*TESTO)
    pdf.set_draw_color(*FILO)
    pdf.set_line_width(0.2)
    pdf.set_fill_color(*RIGA_ALTERNA)
    stile = FontFace(emphasis="BOLD", color=255, fill_color=ARDESIA)
    with pdf.table(first_row_as_headings=True, headings_style=stile, line_height=5.4, text_align="LEFT",
                   col_widths=larghezze, cell_fill_color=RIGA_ALTERNA, cell_fill_mode="ROWS",
                   borders_layout="HORIZONTAL_LINES", padding=1.6) as t:
        r = t.row()
        for h in intestazione:
            r.cell(_latin1(h))
        for riga in righe:
            r = t.row()
            for i, valore in enumerate(riga):
                if i == colonna_priorita:
                    colore, sfondo = PRIORITA.get(str(valore).lower(), (TESTO, RIGA_ALTERNA))
                    r.cell(_latin1(str(valore).upper()), align="C",
                           style=FontFace(emphasis="BOLD", color=colore, fill_color=sfondo))
                else:
                    r.cell(_latin1(str(valore or "-")))
    pdf.ln(2)


def make_verbale_hse(d, foto=()):
    """Verbale di sopralluogo HSE con impaginazione fissa: la grafica la decide il codice, non il modello."""
    pdf = _Report(d.get("titolo") or "Verbale di sopralluogo HSE")
    pdf.add_page()
    nc, azioni = d.get("non_conformita") or [], d.get("azioni") or []

    # Dati del sopralluogo: riquadro con quattro campi
    y = pdf.get_y()
    _box(pdf, 18, y, pdf.epw, 24, RIGA_ALTERNA)
    campi = [("Luogo", d.get("luogo")), ("Data", d.get("data")), ("Ispettore", d.get("ispettore")),
             ("Fonte", "foto inviata dall'utente" if foto else "documenti caricati")]
    for i, (etichetta, valore) in enumerate(campi):
        x, yy = 23 + (i % 2) * (pdf.epw / 2), y + 4 + (i // 2) * 9.5
        pdf.set_xy(x, yy)
        pdf.set_font("Helvetica", size=7.5)
        pdf.set_text_color(*GRIGIO)
        pdf.cell(0, 3.5, etichetta)
        pdf.set_xy(x, yy + 3.6)
        if valore:
            pdf.set_font("Helvetica", "B", 10)
            pdf.set_text_color(*TESTO)
            pdf.cell(pdf.epw / 2 - 10, 5, _latin1(valore))
        else:
            pdf.set_draw_color(*FILO)
            pdf.line(x, yy + 8, x + pdf.epw / 2 - 16, yy + 8)
    pdf.set_y(y + 28)

    # Oggetto e sintesi, con la foto accanto
    _titolo_sezione(pdf, "Oggetto e sintesi")
    y0 = pdf.get_y()
    larghezza_testo = pdf.epw
    if foto:
        dati, proporzione = foto_pronta(foto[0])
        lf = 72
        hf = min(lf / proporzione, 80)
        lf = hf * proporzione
        pdf.image(BytesIO(dati), x=pdf.w - 18 - lf, y=y0, w=lf, h=hf)
        pdf.set_xy(pdf.w - 18 - lf, y0 + hf + 0.5)
        pdf.set_font("Helvetica", "I", 7.5)
        pdf.set_text_color(*GRIGIO)
        pdf.cell(lf, 4, "Foto 1: fornita dall'utente, non verificata sul posto", align="C")
        larghezza_testo = pdf.epw - lf - 6
        fondo_foto = y0 + hf + 5
    else:
        fondo_foto = y0
    pdf.set_xy(18, y0)
    pdf.set_font("Helvetica", size=10)
    pdf.set_text_color(*TESTO)
    pdf.multi_cell(larghezza_testo, 5.6, _latin1(d.get("oggetto", "")), new_x="LMARGIN", new_y="NEXT", align="L")
    pdf.ln(3)
    conteggio = {p: sum(1 for n in nc if str(n.get("priorita", "")).lower() == p) for p in PRIORITA}
    pdf.set_font("Helvetica", "B", 22)
    pdf.set_text_color(*ARDESIA)
    pdf.cell(14, 10, str(len(nc)))
    pdf.set_font("Helvetica", size=10)
    pdf.cell(larghezza_testo - 14, 10, "possibili non conformità", new_x="LMARGIN", new_y="NEXT")
    for p, (colore, sfondo) in PRIORITA.items():
        if conteggio[p]:
            x, y = pdf.get_x(), pdf.get_y() + 1
            _box(pdf, x, y, 30, 7, sfondo)
            pdf.set_xy(x, y)
            pdf.set_font("Helvetica", "B", 8.5)
            pdf.set_text_color(*colore)
            pdf.cell(30, 7, _latin1(f"{conteggio[p]} priorità {p}"), align="C")
            pdf.set_xy(x + 33, y - 1)
    pdf.set_y(max(pdf.get_y() + 12, fondo_foto))

    _titolo_sezione(pdf, "Non conformità")
    _tabella(pdf, ["#", "Osservazione", "Rischio", "Riferimento", "Priorità"],
             [[i, n.get("osservazione"), n.get("rischio"), n.get("riferimento"), n.get("priorita")]
              for i, n in enumerate(nc, 1)], (6, 36, 26, 22, 12), colonna_priorita=4)

    _titolo_sezione(pdf, "Azioni correttive")
    _tabella(pdf, ["#", "Azione", "Rif. NC", "Priorità"],
             [[i, a.get("azione"), a.get("rif_non_conformita"), a.get("priorita")] for i, a in enumerate(azioni, 1)],
             (6, 64, 12, 12), colonna_priorita=3)

    limiti = d.get("limiti") or []
    if limiti:
        _titolo_sezione(pdf, "Limiti della valutazione")
        pdf.set_font("Helvetica", size=9.5)
        testo = _latin1("\n".join(limiti))
        righe = len(pdf.multi_cell(pdf.epw - 12, 5.2, testo, dry_run=True, output="LINES"))
        y = pdf.get_y()
        _box(pdf, 18, y, pdf.epw, righe * 5.2 + 6, (252, 243, 225))
        pdf.set_fill_color(*AMBRA)
        pdf.rect(18, y, 1.8, righe * 5.2 + 6, "F")
        pdf.set_xy(24, y + 3)
        pdf.set_text_color(*TESTO)
        pdf.multi_cell(pdf.epw - 12, 5.2, testo, new_x="LMARGIN", new_y="NEXT", align="L")
        pdf.set_y(y + righe * 5.2 + 9)

    # Firme
    if pdf.get_y() > pdf.h - 50:
        pdf.add_page()
    _titolo_sezione(pdf, "Firme")
    y = pdf.get_y()
    for i, etichetta in enumerate(["Ispettore", "Presa visione (RSPP o responsabile)"]):
        x = 18 + i * (pdf.epw / 2 + 3)
        w = pdf.epw / 2 - 3
        pdf.set_draw_color(*FILO)
        pdf.rect(x, y, w, 24, "D", round_corners=True, corner_radius=2)
        pdf.set_xy(x + 4, y + 3)
        pdf.set_font("Helvetica", size=8)
        pdf.set_text_color(*GRIGIO)
        pdf.cell(w - 8, 4, _latin1(etichetta))
        pdf.line(x + 4, y + 18, x + w - 4, y + 18)
        pdf.set_xy(x + 4, y + 18.5)
        pdf.cell(w - 8, 4, "Firma e data")
    return bytes(pdf.output())


def _e_campo(nodo):
    return isinstance(nodo, dict) and "valore" in nodo and "stato" in nodo


def _testo_valore(v):
    if isinstance(v, bool):  # prima dei numeri: in Python True è anche un intero
        return "sì" if v else "no"
    if v is None or v == "":
        return "-"
    if isinstance(v, float):
        return f"{v:,.3f}".rstrip("0").rstrip(".").replace(",", "X").replace(".", ",").replace("X", ".")
    if isinstance(v, list):
        return "; ".join(_testo_valore(x) for x in v)
    if isinstance(v, dict):
        return "; ".join(f"{k}: {_testo_valore(x)}" for k, x in v.items())
    return str(v)


def _fonte_breve(v):
    """Nome del documento leggibile e spezzabile in tabella: senza estensione, con spazi al posto di _."""
    import re
    testo = _testo_valore(v)
    return re.sub(r"\.(pdf|xlsx|docx)\b", "", testo, flags=re.I).replace("_", " ")


def _titolo_chiave(k):
    return str(k).replace("_", " ").strip().capitalize()


def _dpp_markdown(oggetto):
    """Trasforma il JSON del passaporto in markdown semplice: sezioni, tabelle Campo/Valore/Fonte/Stato, elenchi."""
    righe = []

    def tabella(campi):
        righe.append("| Campo | Valore | Fonte | Stato |")
        for nome, c in campi:
            stato = str(c.get("stato") or "")
            righe.append(f"| {_titolo_chiave(nome)} | {_testo_valore(c.get('valore'))} | "
                         f"{_fonte_breve(c.get('fonte'))} | {stato} |")
        righe.append("")

    def cella(v):
        if _e_campo(v):
            return f"{_testo_valore(v.get('valore'))} ({v.get('stato')})"
        return _testo_valore(v)

    def tabella_record(record):
        """Lista di oggetti simili (es. un componente per riga): una tabella con una colonna per chiave."""
        chiavi = []
        for r in record:
            for k in r:
                if k != "citazione" and k not in chiavi:
                    chiavi.append(k)
        righe.append("| " + " | ".join(_titolo_chiave(k) for k in chiavi) + " |")
        for r in record:
            righe.append("| " + " | ".join(_fonte_breve(r.get(k)) if k == "fonte" else cella(r.get(k))
                                           for k in chiavi) + " |")
        righe.append("")

    def sezione(nodo, livello):
        if _e_campo(nodo):
            tabella([("valore", nodo)])
            return
        if isinstance(nodo, dict):
            campi = [(k, v) for k, v in nodo.items() if _e_campo(v)]
            if campi:
                tabella(campi)
            semplici = [(k, v) for k, v in nodo.items()
                        if not _e_campo(v) and not isinstance(v, (dict, list)) and k != "citazione"]
            if len(semplici) >= 3:  # oggetto con molti valori semplici: tabella a due colonne
                righe.append("| Voce | Valore |")
                for k, v in semplici:
                    righe.append(f"| {_titolo_chiave(k)} | {_testo_valore(v)} |")
                righe.append("")
            for k, v in nodo.items():
                if _e_campo(v) or k == "citazione" or (len(semplici) >= 3 and (k, v) in semplici):
                    continue
                if isinstance(v, (dict, list)):
                    righe.append("#" * min(livello, 3) + " " + _titolo_chiave(k))
                    sezione(v, livello + 1)
                else:
                    righe.append(f"{_titolo_chiave(k)}: {_testo_valore(v)}")
            return
        if isinstance(nodo, list):
            if nodo and all(_e_campo(x) for x in nodo):
                tabella([(x.get("campo") or x.get("nome") or f"voce {i + 1}", x) for i, x in enumerate(nodo)])
            elif nodo and all(isinstance(x, dict) for x in nodo) and all(
                    all(_e_campo(v) or not isinstance(v, (dict, list)) for v in x.values()) for x in nodo):
                tabella_record(nodo)
            else:
                for i, x in enumerate(nodo):
                    if isinstance(x, dict):
                        nome = x.get("componente") or x.get("nome") or x.get("campo") or f"Voce {i + 1}"
                        righe.append("#" * min(livello, 3) + " " + _testo_valore(nome))
                        sezione({k: v for k, v in x.items() if k not in ("componente", "nome")}, livello + 1)
                    else:
                        righe.append(f"- {_testo_valore(x)}")
            righe.append("")
            return
        righe.append(_testo_valore(nodo))

    for k, v in oggetto.items():
        righe.append("## " + _titolo_chiave(k))
        sezione(v, 3)
        righe.append("")
    return "\n".join(righe)


def make_dpp_pdf(oggetto, nome_file, titolo, grafico=None):
    """PDF del passaporto generato dal codice, a partire dal JSON già verificato."""
    contenuto = ("Documento dimostrativo generato dal codice a partire dal JSON verificato. "
                 "I requisiti di dettaglio per l'acciaio dipendono dall'atto delegato ESPR.\n\n" + _dpp_markdown(oggetto))
    return make_document("pdf", nome_file, titolo, contenuto, grafico)


def make_document(formato, nome_file, titolo, contenuto, grafico=None, foto=None):
    """Restituisce (bytes, nome del file, tipo MIME)."""
    formato = formato.lower()
    if formato == "xlsx":
        dati = _xlsx(titolo, contenuto)
    elif formato == "docx":
        dati = _docx(titolo, contenuto, grafico, foto)
    elif formato == "pdf":
        dati = _pdf(titolo, contenuto, grafico, foto)
    elif formato == "json":
        testo = contenuto.strip().removeprefix("```json").removeprefix("```").removesuffix("```")
        try:
            oggetto = json.loads(testo)
        except json.JSONDecodeError as errore:
            raise ValueError(f"il contenuto non è JSON valido ({errore}): correggilo e riprova") from errore
        dati = json.dumps(oggetto, ensure_ascii=False, indent=2).encode("utf-8")
    else:
        raise ValueError(f"Formato non supportato: {formato}")
    nome = re.sub(r"[^\w\-]+", "_", nome_file).strip("_") or "documento"
    return dati, f"{nome}.{formato}", MIME[formato]
